import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def markdown_to_docx(md_path, docx_path):
    doc = docx.Document()

    # Set Margins to 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set Normal Style Font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_buffer = []
    in_table = False
    table_data = []

    def flush_table(doc, table_data):
        if not table_data:
            return
        
        cols = max(len(row) for row in table_data)
        table = doc.add_table(rows=len(table_data), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row in enumerate(table_data):
            is_header = (r_idx == 0)
            for c_idx, val in enumerate(row):
                if c_idx < len(table.rows[r_idx].cells):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = val.strip()
                    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    
                    if is_header:
                        set_cell_background(cell, "1E3A8A") # Deep Blue
                        for run in p.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.font.size = Pt(10)
                    else:
                        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_color)
                        for run in p.runs:
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        doc.add_paragraph() # Spacing after table

    i = 0
    while i < len(lines):
        line = lines[i]
        raw_line = line.rstrip('\n')

        # Table processing
        if '|' in raw_line and not in_code_block:
            parts = [p.strip() for p in raw_line.split('|')[1:-1]]
            # Check if separator row like |---|---|
            if parts and all(re.match(r'^:?-+:?$', p) for p in parts):
                i += 1
                continue
            if parts:
                if not in_table:
                    in_table = True
                    table_data = []
                table_data.append(parts)
                i += 1
                continue

        if in_table and ('|' not in raw_line or in_code_block):
            in_table = False
            flush_table(doc, table_data)
            table_data = []

        # Code block processing
        if raw_line.startswith('```'):
            if in_code_block:
                # Flush code block
                code_text = '\n'.join(code_buffer)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
                code_buffer = []
            i += 1
            continue

        if in_code_block:
            code_buffer.append(raw_line)
            i += 1
            continue

        # Blank lines
        if not raw_line.strip():
            i += 1
            continue

        # Headings
        if raw_line.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            run = h.add_run(raw_line[2:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep navy
        elif raw_line.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(raw_line[3:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue accent
        elif raw_line.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(raw_line[4:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        elif raw_line.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(raw_line[5:].strip())
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        # Blockquotes / Alerts
        elif raw_line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            text = raw_line[2:].strip()
            # Clean markdown inline bold/italic/links
            text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
            text = text.replace('**', '').replace('*', '').replace('`', '')
            run = p.add_run(text)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        # Bullet points
        elif raw_line.startswith('* ') or raw_line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            text = raw_line[2:].strip()
            # Inline formatting
            parse_inline_formatting(p, text)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_inline_formatting(p, raw_line)

        i += 1

    if in_table:
        flush_table(doc, table_data)

    doc.save(docx_path)
    print(f"Document successfully saved to {docx_path}")

def parse_inline_formatting(p, text):

    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.font.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.font.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        else:
            p.add_run(part)

if __name__ == '__main__':
    md_file = r"C:\Users\adipk\.gemini\antigravity\brain\70d341e0-ed21-41f2-867d-aa04855e36d0\phishguard_comprehensive_technical_report.md"
    os.makedirs(r"c:\Users\adipk\Desktop\PhishGuard\docs", exist_ok=True)
    docx_file = r"c:\Users\adipk\Desktop\PhishGuard\docs\PhishGuard_Q1_Journal_Research_Paper.docx"
    markdown_to_docx(md_file, docx_file)
    
    # Also save a Markdown copy in docs
    md_copy = r"c:\Users\adipk\Desktop\PhishGuard\docs\PhishGuard_Q1_Journal_Research_Paper.md"
    with open(md_file, 'r', encoding='utf-8') as src, open(md_copy, 'w', encoding='utf-8') as dst:
        dst.write(src.read())
    print(f"Markdown copy saved to {md_copy}")
